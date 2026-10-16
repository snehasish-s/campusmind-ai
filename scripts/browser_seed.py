"""Seed a dedicated browser-test institution; password arrives via process environment."""

import os
import sys
from pathlib import Path
from datetime import timedelta

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "services/api"))
from docx import Document  # noqa: E402
from app.db import SessionLocal  # noqa: E402
from app.models import Institution, Department, User, StudentProfile, Company, RecruitmentDrive, ActivityEvent, now  # noqa: E402
from app.security import password_hash  # noqa: E402


def seed():
    password = os.environ["CM_E2E_PASSWORD"]
    slug = os.environ["CM_E2E_INSTITUTION"]
    fixture = Document()
    fixture.add_heading("Synthetic browser test resume", 0)
    fixture.add_paragraph("Python, SQL and accessible web application development.")
    fixture.save(Path(__file__).resolve().parents[1] / "data/browser-resume.docx")
    with SessionLocal() as db:
        institution = Institution(slug=slug, name="Synthetic Browser Test Institute")
        db.add(institution)
        db.flush()
        departments = []
        for code in ["CSE", "ECE", "EE", "ME", "CE"]:
            department = Department(institution_id=institution.id, code=code, name=f"Synthetic {code} Department")
            db.add(department)
            db.flush()
            departments.append(department)
        for login, role in [
            ("ADMIN", "institution_admin"),
            ("PLACEMENT", "placement_officer"),
            ("STUDENT", "student"),
            ("JUNIOR", "student"),
            ("TEACHER", "faculty"),
        ]:
            user = User(
                institution_id=institution.id,
                roll_no=login,
                email=f"{login.lower()}@example.com",
                full_name=f"Synthetic {login.title()}",
                role=role,
                active=True,
                password_hash=password_hash(password),
                department_id=departments[0].id if role in {"student", "faculty"} else None,
            )
            db.add(user)
            db.flush()
            if role == "student":
                db.add(
                    StudentProfile(
                        institution_id=institution.id,
                        user_id=user.id,
                        graduation_year=2027,
                        cgpa=8.5,
                        academic_year=2 if login == "JUNIOR" else 4,
                    )
                )
        for i in range(25):
            user = User(
                institution_id=institution.id,
                roll_no=f"SYN{i:05d}",
                email=f"synthetic{i}@example.com",
                full_name=f"Synthetic Student {i:03d}",
                role="student",
                department_id=departments[i % 5].id,
            )
            db.add(user)
            db.flush()
            db.add(StudentProfile(institution_id=institution.id, user_id=user.id, graduation_year=2027, cgpa=7 + i % 20 / 10))
        company = Company(institution_id=institution.id, name="Northstar Labs (Synthetic)", sector="Software")
        db.add(company)
        db.flush()
        db.add(
            RecruitmentDrive(
                institution_id=institution.id,
                company_id=company.id,
                position="Graduate Software Engineer",
                description="Synthetic recruitment opportunity for testing the student placement workflow.",
                employment_type="Full-time",
                location="Bhubaneswar / Hybrid",
                ctc="Illustrative ₹5 LPA",
                min_cgpa=6,
                max_backlogs=0,
                eligible_branches=["CSE", "ECE"],
                graduation_year=2027,
                required_skills=["Python", "SQL"],
                deadline=now() + timedelta(days=30),
                drive_date=now() + timedelta(days=35),
                selection_rounds=["Technical interview", "HR conversation"],
                status="published",
            )
        )
        db.add(ActivityEvent(institution_id=institution.id, title="Synthetic test workspace initialized", kind="institution"))
        db.commit()


if __name__ == "__main__":
    seed()
