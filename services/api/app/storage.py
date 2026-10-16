from io import BytesIO
from pathlib import Path
from typing import Protocol
import uuid
import zipfile
import httpx
from urllib.parse import quote
from fastapi import HTTPException
from fastapi.responses import Response, FileResponse
from pypdf import PdfReader
from docx import Document
from .config import settings

MAX_FILE_SIZE = 8 * 1024 * 1024


class Storage(Protocol):
    def put(self, tenant: str, data: bytes, suffix: str) -> str: ...
    def path(self, key: str) -> Path: ...
    def delete(self, key: str) -> None: ...


class LocalStorage:
    def __init__(self):
        self.root = Path(settings().storage_root).resolve()

    def put(self, tenant, data, suffix):
        key = f"{tenant}/{uuid.uuid4().hex}{suffix}"
        path = self.path(key)
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(data)
        try:
            path.chmod(0o600)
        except OSError:
            pass
        return key

    def path(self, key):
        path = (self.root / key).resolve()
        if not path.is_relative_to(self.root):
            raise ValueError("Invalid storage key")
        return path

    def delete(self, key):
        self.path(key).unlink(missing_ok=True)

    def read(self, key):
        return self.path(key).read_bytes()


class SupabaseStorage:
    """Private bucket, backend-only credential. Objects are never made public."""

    def __init__(self):
        config = settings()
        if not config.supabase_url or not config.supabase_service_role_key:
            raise HTTPException(503, "Private storage credentials have not been configured.")
        self.url = config.supabase_url.rstrip("/") + "/storage/v1/object"
        self.bucket = config.supabase_storage_bucket
        self.headers = {"Authorization": "Bearer " + config.supabase_service_role_key, "apikey": config.supabase_service_role_key}

    def safe_key(self, key):
        if ".." in key or key.startswith("/") or "\\" in key:
            raise ValueError("Invalid storage key")
        return quote(self.bucket + "/" + key, safe="/")

    def put(self, tenant, data, suffix):
        key = f"{tenant}/{uuid.uuid4().hex}{suffix}"
        try:
            response = httpx.post(
                self.url + "/" + self.safe_key(key),
                headers={**self.headers, "Content-Type": "application/octet-stream", "x-upsert": "false"},
                content=data,
                timeout=30,
            )
            response.raise_for_status()
        except httpx.HTTPError:
            raise HTTPException(503, "Private object storage is unavailable.")
        return key

    def read(self, key):
        try:
            response = httpx.get(self.url + "/authenticated/" + self.safe_key(key), headers=self.headers, timeout=30)
            response.raise_for_status()
            return response.content
        except httpx.HTTPError:
            raise HTTPException(503, "The private document could not be retrieved.")

    def delete(self, key):
        self.safe_key(key)
        try:
            response = httpx.request(
                "DELETE", self.url + "/" + quote(self.bucket, safe=""), headers=self.headers, json={"prefixes": [key]}, timeout=30
            )
            response.raise_for_status()
        except httpx.HTTPError:
            raise HTTPException(503, "Private storage deletion could not be completed. Please retry.")


def get_storage():
    return SupabaseStorage() if settings().storage_backend == "supabase" else LocalStorage()


def document_response(key, filename, mime):
    storage = get_storage()
    if isinstance(storage, LocalStorage):
        return FileResponse(
            storage.path(key),
            filename=filename,
            media_type=mime,
            headers={"Cache-Control": "no-store", "X-Content-Type-Options": "nosniff"},
        )
    return Response(
        storage.read(key),
        media_type=mime,
        headers={
            "Content-Disposition": "attachment; filename*=UTF-8''" + quote(filename, safe=""),
            "Cache-Control": "no-store",
            "X-Content-Type-Options": "nosniff",
        },
    )


def validate_document(filename, content, allow_text=False):
    if not content or len(content) > MAX_FILE_SIZE:
        raise HTTPException(422, "Upload a non-empty file no larger than 8 MB.")
    suffix = Path(filename).suffix.lower()
    try:
        if suffix == ".txt" and allow_text:
            text = content.decode("utf-8-sig")
            if "\x00" in text:
                raise ValueError("Binary content is not text")
            mime = "text/plain"
        elif suffix == ".pdf":
            if not content.startswith(b"%PDF-"):
                raise ValueError("Invalid PDF signature")
            reader = PdfReader(BytesIO(content), strict=True)
            if reader.is_encrypted or len(reader.pages) > 100:
                raise ValueError("Encrypted or oversized document")
            root = reader.trailer["/Root"]
            if root.get("/OpenAction") or root.get("/AA"):
                raise ValueError("Active PDF content is not supported")
            for page in reader.pages:
                if page.get("/AA"):
                    raise ValueError("Active PDF content is not supported")
            text = "\n".join((page.extract_text() or "") for page in reader.pages)
            mime = "application/pdf"
        elif suffix == ".docx":
            with zipfile.ZipFile(BytesIO(content)) as archive:
                entries = archive.infolist()
                if len(entries) > 1500 or sum(e.file_size for e in entries) > 32 * 1024 * 1024:
                    raise ValueError("Archive exceeds limits")
                names = archive.namelist()
                if "word/document.xml" not in names or "[Content_Types].xml" not in names:
                    raise ValueError("Invalid DOCX structure")
                if any("vbaproject" in n.lower() or "/embeddings/" in n.lower() for n in names):
                    raise ValueError("Embedded executable content is not supported")
            document = Document(BytesIO(content))
            text = "\n".join(
                [p.text for p in document.paragraphs] + [c.text for table in document.tables for row in table.rows for c in row.cells]
            )
            mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        else:
            raise ValueError("Unsupported format")
    except Exception:
        raise HTTPException(422, "Invalid or unsupported document. Use an unencrypted PDF or DOCX without active content.")
    if len(text) > 1_000_000:
        raise HTTPException(422, "Document text exceeds the extraction limit.")
    return suffix, mime, text
