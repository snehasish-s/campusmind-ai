import csv
import io
import secrets
from datetime import timedelta
from docx import Document
from sqlalchemy import select, func
from fastapi.testclient import TestClient
from app.main import app
from app.db import SessionLocal
from app.models import User, Invitation, ImportJob, Resume, RecruitmentDrive, now
from app.security import digest
from conftest import login, secure


def csv_data(count=1000, cgpa=8.2):
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["roll_no", "full_name", "email", "department", "graduation_year", "cgpa", "active_backlogs", "phone", "academic_year"])
    for i in range(count):
        writer.writerow([f"SYN{i:05d}", f"Synthetic Student {i}", f"student{i}@example.com", "CSE", 2027, cgpa, 0, "", 4])
    return output.getvalue()


def preview(client, content, policy="skip"):
    result = client.post(
        "/api/imports/preview", files={"file": ("synthetic.csv", content, "text/csv")}, data={"policy": policy}, headers=secure(client)
    )
    assert result.status_code == 200, result.text
    return result.json()


def import_students(client, count=1):
    job = preview(client, csv_data(count))
    result = client.post(f"/api/imports/{job['id']}/confirm", headers=secure(client))
    assert result.status_code == 200, result.text
    return result.json()


def document_bytes():
    document = Document()
    document.add_heading("Synthetic Student", 0)
    document.add_paragraph("Python, SQL and a synthetic campus placement project.")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def drive_payload():
    return {
        "company_name": "Synthetic Labs",
        "position": "Graduate Engineer",
        "description": "Build and test software in a synthetic demonstration role.",
        "employment_type": "Full-time",
        "location": "Bhubaneswar",
        "ctc": "Illustrative ₹5 LPA",
        "required_skills": ["Python"],
        "min_cgpa": 6,
        "max_backlogs": 0,
        "eligible_branches": ["CSE"],
        "graduation_year": 2027,
        "deadline": (now() + timedelta(days=10)).isoformat(),
        "drive_date": (now() + timedelta(days=15)).isoformat(),
        "selection_rounds": ["Technical interview"],
        "status": "published",
    }


def test_complete_1000_student_acceptance_flow(setup):
    officer, password, mail, ids = setup
    # 1. Institution admin authenticates; protected cookies are HTTP-only.
    login(officer, password, identifier="ADMIN")
    assert officer.get("/api/staff").status_code == 200
    login(officer, password)
    # 2-3. Preview and persist 1,000 students through the real import service.
    job = import_students(officer, 1000)
    assert job["status"] == "completed" and job["created"] == 1000 and job["errors"] == 0
    assert len(mail) == 1000
    assert officer.get("/api/students").json()["total"] == 1000
    assert officer.get("/api/dashboard").json()["students"] == 1000
    assert officer.post(f"/api/imports/{job['id']}/confirm", headers=secure(officer)).json()["created"] == 1000
    assert preview(officer, csv_data())["id"] == job["id"]
    with SessionLocal() as db:
        target = db.scalar(select(User).where(User.roll_no == "SYN00000"))
        student_id = target.id
        assert target.password_hash is None
        assert db.scalar(select(func.count()).select_from(ImportJob)) == 1
    # An audited manual link rotates old links; raw invitation tokens are never stored.
    distribution = officer.post(f"/api/auth/invitations/{student_id}/distribute", headers=secure(officer)).json()
    token = distribution["activation_url"].split("#token=")[1]
    with SessionLocal() as db:
        invitation = db.scalar(select(Invitation).where(Invitation.token_hash == digest(token)))
        assert invitation and invitation.token_hash != token
    # 4-5. Student activation, single-use enforcement, login and profile editing.
    student = TestClient(app)
    chosen = secrets.token_urlsafe(20)
    activation = {"token": token, "password": chosen, "consent": True}
    assert student.post("/api/auth/activate", json=activation).status_code == 200
    assert student.post("/api/auth/activate", json=activation).status_code == 400
    login(student, chosen, identifier="SYN00000")
    profile = {
        "full_name": "Synthetic Student Updated",
        "graduation_year": 2027,
        "cgpa": 8.7,
        "active_backlogs": 0,
        "phone": "",
        "bio": "Synthetic test profile",
        "projects": [{"title": "Campus project", "description": "Synthetic demonstration"}],
        "internships": [],
    }
    assert student.put("/api/profile", json=profile, headers=secure(student)).status_code == 200
    assert student.get("/api/profile").json()["cgpa"] == 8.7
    # 6. DOCX upload, text extraction and version history.
    upload = student.post("/api/resumes", files={"file": ("resume.docx", document_bytes())}, headers=secure(student))
    assert upload.status_code == 200, upload.text
    resume_id = upload.json()["id"]
    assert "Python" in student.get(f"/api/resumes/{resume_id}/text").json()["text"]
    second = student.post("/api/resumes", files={"file": ("updated.docx", document_bytes())}, headers=secure(student))
    assert second.json()["version"] == 2
    # 7. Authorized officer access.
    assert officer.get(f"/api/students/{student_id}").json()["cgpa"] == 8.7
    assert officer.get(f"/api/resumes/{resume_id}/download").status_code == 200
    # 8-9. Publish a real drive, browse it, register and see status tracking.
    drive = officer.post("/api/drives", json=drive_payload(), headers=secure(officer))
    assert drive.status_code == 200, drive.text
    drive_id = drive.json()["id"]
    assert student.get("/api/drives").json()[0]["id"] == drive_id
    application = student.post(f"/api/drives/{drive_id}/apply", headers=secure(student))
    assert application.status_code == 200
    assert student.post(f"/api/drives/{drive_id}/apply", headers=secure(student)).json()["id"] == application.json()["id"]
    assert student.get("/api/applications").json()[0]["status"] == "registered"
    # 10. Cross-college access to student, file, import and drive mutation is denied.
    other = TestClient(app)
    login(other, password, institution="beta")
    for path in [
        f"/api/students/{student_id}",
        f"/api/resumes/{resume_id}/download",
        f"/api/resumes/{resume_id}/text",
        f"/api/imports/{job['id']}",
        f"/api/imports/{job['id']}/errors.csv",
    ]:
        assert other.get(path).status_code == 404, path
    assert other.put(f"/api/drives/{drive_id}", json=drive_payload(), headers=secure(other)).status_code == 404
    assert other.get("/api/drives").json() == []
    assert other.get("/api/students").json()["total"] == 0
    with SessionLocal() as db:
        assert db.scalar(select(func.count()).select_from(Resume)) == 2
        assert db.scalar(select(func.count()).select_from(RecruitmentDrive)) == 1
    student.close()
    other.close()


def test_csv_validation_mapping_duplicate_policy_and_idempotency(setup):
    client, password, mail, ids = setup
    login(client, password)
    content = csv_data(2) + "SYN00000,Duplicate,duplicate@example.com,CSE,2027,8,0,\nBAD,Invalid,not-an-email,CSE,2027,99,0,\n"
    job = preview(client, content)
    assert job["errors"] == 2 and job["total"] == 4
    assert client.get(f"/api/imports/{job['id']}/errors.csv").text.count("\n") == 3
    result = client.post(f"/api/imports/{job['id']}/confirm", headers=secure(client)).json()
    assert result["created"] == 2
    skip = preview(client, csv_data(2, cgpa=7.1))
    assert client.post(f"/api/imports/{skip['id']}/confirm", headers=secure(client)).json()["skipped"] == 2
    update = preview(client, csv_data(2, cgpa=7.1), "update")
    assert client.post(f"/api/imports/{update['id']}/confirm", headers=secure(client)).json()["updated"] == 2
    assert client.get("/api/students").json()["items"][0]["cgpa"] == 7.1
    mapped = csv_data(1).replace("roll_no", "Student ID").replace("SYN00000", "MAPPED01").replace("student0@", "mapped0@")
    response = client.post(
        "/api/imports/preview", files={"file": ("mapped.csv", mapped)}, data={"mapping": '{"roll_no":"Student ID"}'}, headers=secure(client)
    )
    assert response.status_code == 200 and response.json()["errors"] == 0


def test_csrf_rbac_expiry_and_upload_validation(setup):
    client, password, mail, ids = setup
    assert client.get("/api/students").status_code == 401
    login(client, password)
    assert client.post("/api/drives", json=drive_payload()).status_code == 403
    assert (
        client.post("/api/drives", json=drive_payload(), headers={**secure(client), "Origin": "https://untrusted.example"}).status_code
        == 403
    )
    assert client.post("/api/departments", json={"code": "EE", "name": "Electrical"}, headers=secure(client)).status_code == 403
    import_students(client)
    student_id = client.get("/api/students").json()["items"][0]["id"]
    link = client.post(f"/api/auth/invitations/{student_id}/distribute", headers=secure(client)).json()["activation_url"]
    token = link.split("#token=")[1]
    assert client.post(f"/api/auth/invitations/{student_id}/revoke", headers=secure(client)).status_code == 200
    assert (
        client.post("/api/auth/activate", json={"token": token, "password": secrets.token_urlsafe(20), "consent": True}).status_code == 400
    )
    link = client.post(f"/api/auth/invitations/{student_id}/distribute", headers=secure(client)).json()["activation_url"]
    chosen = secrets.token_urlsafe(20)
    student = TestClient(app)
    assert (
        student.post("/api/auth/activate", json={"token": link.split("#token=")[1], "password": chosen, "consent": True}).status_code == 200
    )
    login(student, chosen, identifier="SYN00000")
    assert student.get("/api/students").status_code == 403
    assert student.get("/api/staff").status_code == 403
    assert student.get("/api/imports").status_code == 403
    assert student.get("/api/documents").status_code == 403
    assert student.post("/api/drives", json=drive_payload(), headers=secure(student)).status_code == 403
    assert student.post("/api/resumes", files={"file": ("bad.pdf", b"not a pdf")}, headers=secure(student)).status_code == 422
    assert student.post("/api/resumes", files={"file": ("program.exe", b"MZ")}, headers=secure(student)).status_code == 422
    assert (
        student.post("/api/resumes", files={"file": ("huge.pdf", b"x" * (8 * 1024 * 1024 + 1))}, headers=secure(student)).status_code == 422
    )
    assert student.post("/api/auth/logout", headers=secure(student)).status_code == 200
    assert student.get("/api/profile").status_code == 401
    student.close()


def test_invitation_expiration_reset_rate_limits_and_session_revocation(setup):
    client, password, mail, ids = setup
    login(client, password)
    import_students(client)
    target = client.get("/api/students").json()["items"][0]["id"]
    token = client.post(f"/api/auth/invitations/{target}/distribute", headers=secure(client)).json()["activation_url"].split("#token=")[1]
    with SessionLocal() as db:
        inv = db.scalar(select(Invitation).where(Invitation.token_hash == digest(token)))
        inv.expires_at = now() - timedelta(seconds=1)
        db.commit()
    assert (
        client.post("/api/auth/activate", json={"token": token, "password": secrets.token_urlsafe(20), "consent": True}).status_code == 400
    )
    reset = client.post("/api/auth/password-reset", json={"institution": "alpha", "identifier": "OFFICER"})
    assert reset.status_code == 200
    reset_token = mail[-1].get_content().split("#token=")[1].split()[0]
    new_password = secrets.token_urlsafe(20)
    assert client.post("/api/auth/activate", json={"token": reset_token, "password": new_password}).status_code == 200
    assert client.get("/api/auth/me").status_code == 401
    login(client, new_password)
    for _ in range(11):
        result = client.post("/api/auth/login", json={"institution": "alpha", "identifier": "UNKNOWN", "password": "incorrect"})
    assert result.status_code == 429
